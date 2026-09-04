"""
Text extraction from PDF (native + scanned), DOCX, XLSX, plain text.

Uses PyMuPDF (fitz) — pip-installable, no system-level PDF library needed.
For scanned pages (where PyMuPDF returns almost no text), we rasterise the
page at 200 dpi and hand it to Tesseract for OCR.
"""
import io
from pathlib import Path

import docx
import fitz  # type: ignore[import-untyped] # PyMuPDF
import openpyxl  # type: ignore
import pytesseract
from PIL import Image


def extract_text(path: Path) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext in (".txt", ".md"):
        return path.read_text(errors="ignore")
    raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(path: Path) -> str:
    """Page-by-page PDF text with OCR fallback for scanned pages."""
    doc = fitz.open(path)
    out = []
    for i, page in enumerate(doc, 1):
        text = page.get_text().strip()
        if len(text) < 40:
            # Very little native text → probably a scanned page. OCR it.
            pix = page.get_pixmap(dpi=200)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(img).strip()
        out.append(f"--- Page {i} ---\n{text}")
    doc.close()
    return "\n\n".join(out)


def _extract_docx(path: Path) -> str:
    """Grab all non-empty paragraphs from a Word document."""
    d = docx.Document(str(path))
    return "\n\n".join(p.text for p in d.paragraphs if p.text.strip())


def _extract_xlsx(path: Path) -> str:
    """Flatten every sheet's populated cells into a text stream."""
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)